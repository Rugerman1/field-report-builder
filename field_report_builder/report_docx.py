"""Word export (python-docx), so the report can be edited before it is signed."""

from __future__ import annotations

import io

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from . import __version__, layout
from .model import Report
from .photos import jpeg_bytes

TEXT_WIDTH = 7.3  # inches: letter with 0.6 in margins
DARK = "2F3B4C"
SHADE = "E8EBEF"


def _shade(cell, hex_color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tc_pr.append(shd)


def _cell_text(cell, text, bold=False, size=9, color=None):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(text or "")
    run.bold = bold
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def _widths(table, fractions):
    for row in table.rows:
        for cell, f in zip(row.cells, fractions):
            cell.width = Inches(TEXT_WIDTH * f)


def _bar(doc, title):
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    c = t.rows[0].cells[0]
    _cell_text(c, title.upper(), bold=True, size=10, color="FFFFFF")
    _shade(c, DARK)
    _widths(t, [1.0])
    _gap(doc, 2)


def _gap(doc, pts=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.line_spacing = Pt(pts)


def _para(doc, text, size=9, italic=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(text)
    run.font.size = Pt(size)
    run.italic = italic
    return p


def _table(doc, sec: layout.Table):
    _bar(doc, sec.title)
    if not sec.rows:
        _para(doc, sec.empty)
        _gap(doc)
        return
    t = doc.add_table(rows=1 + len(sec.rows), cols=len(sec.headers))
    t.style = "Table Grid"
    for i, h in enumerate(sec.headers):
        _cell_text(t.rows[0].cells[i], h, bold=True)
        _shade(t.rows[0].cells[i], SHADE)
    for r, row in enumerate(sec.rows, start=1):
        for i, v in enumerate(row):
            _cell_text(t.rows[r].cells[i], v)
    _widths(t, sec.widths)
    _gap(doc)


def _bullets(doc, sec: layout.Bullets):
    _bar(doc, sec.title)
    if not sec.items:
        _para(doc, sec.empty)
    for item in sec.items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(1)
        p.add_run(item).font.size = Pt(9)
    _gap(doc)


def _page_number(paragraph):
    run = paragraph.add_run()
    for kind, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if kind:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        run._r.append(el)
    run.font.size = Pt(7.5)


def build_docx(r: Report, path: str) -> None:
    doc = Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.PORTRAIT
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, side, Inches(0.6))
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(9)

    foot = sec.footer.paragraphs[0]
    foot.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = foot.add_run(f"{r.project}  |  {r.report_date}  |  Field Report Builder {__version__}  |  Page ")
    run.font.size = Pt(7.5)
    _page_number(foot)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    tr = title.add_run(layout.TITLE)
    tr.bold = True
    tr.font.size = Pt(15)

    head = layout.header_rows(r)
    t = doc.add_table(rows=len(head), cols=4)
    t.style = "Table Grid"
    for i, pair in enumerate(head):
        for j, (label, value) in enumerate(pair):
            _cell_text(t.rows[i].cells[2 * j], label.upper(), bold=True, size=7.5, color="444444")
            _shade(t.rows[i].cells[2 * j], SHADE)
            _cell_text(t.rows[i].cells[2 * j + 1], value)
    _widths(t, [0.15, 0.35, 0.15, 0.35])
    _gap(doc, 10)

    _table(doc, layout.weather_table(r))
    note = layout.weather_source_note(r)
    if note:
        _para(doc, note, size=7.5, italic=True)
    for s in layout.sections(r):
        _table(doc, s) if isinstance(s, layout.Table) else _bullets(doc, s)

    _bar(doc, "Certification")
    _para(doc, "I certify that this report is a true record of the work and conditions observed.")
    sig = doc.add_table(rows=1, cols=2)
    sig.style = "Table Grid"
    for i, (label, _) in enumerate(layout.SIGNATURES):
        c = sig.rows[0].cells[i]
        _cell_text(c, label.upper(), bold=True, size=7.5, color="444444")
        c.add_paragraph()
        c.add_paragraph()
        p = c.add_paragraph()
        p.add_run("_" * 34 + "   " + "_" * 12).font.size = Pt(8)
        p = c.add_paragraph()
        p.add_run("Signature" + " " * 52 + "Date").font.size = Pt(7.5)
    _widths(sig, [0.5, 0.5])

    photos = layout.report_photos(r)
    if photos:
        doc.add_page_break()
        _bar(doc, "Photographs")
        grid = doc.add_table(rows=(len(photos) + 1) // 2, cols=2)
        for k, (_, photo, line) in enumerate(photos):
            cell = grid.rows[k // 2].cells[k % 2]
            cell.text = ""
            p = cell.paragraphs[0]
            try:
                p.add_run().add_picture(io.BytesIO(jpeg_bytes(photo.path, max_px=1400)), width=Inches(3.45))
            except OSError:
                p.add_run("(image file missing)")
            cap = cell.add_paragraph()
            cap.add_run(line).font.size = Pt(8.5)
        _widths(grid, [0.5, 0.5])

    doc.core_properties.title = f"Daily Work Report {r.report_date}"
    doc.core_properties.author = r.inspector or "Field Report Builder"
    doc.save(path)
