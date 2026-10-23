from docx import Document
from pypdf import PdfReader

from field_report_builder.extract import fill_report
from field_report_builder.model import Report, Weather
from field_report_builder.photos import caption_photos, read_photo
from field_report_builder.report_docx import build_docx
from field_report_builder.report_pdf import build_pdf
from field_report_builder.sample import example_report, make_sample_photos


def _report(tmp_path):
    r = example_report()
    fill_report(r)
    r.weather = Weather("Clear", "Partly cloudy", 78, 52, 55, 71, 76, 0.0, 9, r.weather.observed, "Open-Meteo")
    r.photos = [read_photo(p) for p in make_sample_photos(str(tmp_path / "photos"))]
    caption_photos(r.photos, r.transcript)
    return r


def test_pdf_has_every_section_and_the_photos(tmp_path):
    path = tmp_path / "r.pdf"
    build_pdf(_report(tmp_path), str(path))
    pdf = PdfReader(str(path))
    text = "\n".join(p.extract_text() for p in pdf.pages)
    for needle in ("DAILY WORK REPORT", "SC 93 Widening", "Thursday, October 15, 2026", "CONTRACTOR FORCES",
                   "Upstate Concrete", "CAT 320 excavator", "Sta. 12+50 to 14+00", "MATERIALS RECEIVED",
                   "TESTS AND INSPECTIONS", "DELAYS, PROBLEMS AND RFIS", "SAFETY", "VISITORS", "CERTIFICATION",
                   "PHOTOGRAPHS", "Photo 1. 10:42 AM", "Photo 2. 2:18 PM", "78 / 52 F"):
        assert needle in text, needle
    assert len(pdf.pages) == 3
    assert sum(len(p.images) for p in pdf.pages) == 2


def test_docx_has_every_section_and_the_photos(tmp_path):
    path = tmp_path / "r.docx"
    build_docx(_report(tmp_path), str(path))
    doc = Document(str(path))
    cells = [c.text for t in doc.tables for row in t.rows for c in row.cells]
    paras = [p.text for p in doc.paragraphs]
    text = "\n".join(cells + paras)
    for needle in ("DAILY WORK REPORT", "P-0293-24", "Palmetto Site Works", "Smooth drum roller",
                   "150 linear feet", "Received six loads of number 57 stone.", "INSPECTOR", "Photo 2. 2:18 PM"):
        assert needle in text, needle
    assert len(doc.inline_shapes) == 2


def test_empty_report_still_exports(tmp_path):
    r = Report(project="Blank")
    build_pdf(r, str(tmp_path / "e.pdf"))
    build_docx(r, str(tmp_path / "e.docx"))
    text = PdfReader(str(tmp_path / "e.pdf")).pages[0].extract_text()
    assert "None reported." in text and "Not recorded." in text


def test_save_and_load_round_trip(tmp_path):
    r = _report(tmp_path)
    path = tmp_path / "day.frb.json"
    r.save(str(path))
    back = Report.load(str(path))
    assert back == r
